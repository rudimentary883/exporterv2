"""
Tabbycat Adjudicator Tab Exporter v1.1
  * v1.1: three export modes (export_mode)
        combined  = the original sheet (both averages)
        feedback  = like the Davao sheet: score + number of feedback per round, average weighted by number of feedback
        rounds    = like the MPDC sheet: one average per round, final average = average of the rounds
                    (+ the Bali UADC layout with "Average with Min. Score Dropped" / "No. of Rounds Judged" / "Dropped Score"
                    when the tournament is UADC)
    Breaking is the 4th column (after Institution) and is written as TRUE/FALSE (Insert > Checkbox turns it into boxes).
    Per-round averages are stored with 4 decimals (shown with 2); the test score is the exact Base Score from Tabbycat.

Tabbycat Adjudicator Tab Exporter v1.0

Connects to a Tabbycat site through its REST API (same token / session-auth
approach as the Tabbycat Importer) and builds an Excel "Adjudicator Tab":

  * average feedback score per round (always 2 decimals)
  * number of feedback received per round
  * score cell colour = role that round
        chair     -> #b10095
        panellist -> #36b6c1
        trainee   -> #e96e20
  * test score, overall average (weighted by number of feedback),
    total feedback, final rank (CAP for adjudication core), breaking

READ-ONLY: this app only ever sends GET requests to Tabbycat
(the only POST is the Django login form used for the optional session fallback).

Admin access is required, because feedback and pairings are not public.
"""

import os
import io
import re
import uuid
from collections import OrderedDict, defaultdict

import requests
from flask import (Flask, render_template, request, send_file, flash,
                   redirect, url_for, jsonify)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'tabbycat-exporter-key-2026')
app.config['MAX_CONTENT_LENGTH'] = 1 * 1024 * 1024

# -----------------------------------------------------------------------------
# Look & feel of the Excel file (change here if you want different colours)
# -----------------------------------------------------------------------------
ROLE_COLORS = {
    'chair': 'B10095',
    'panellist': '36B6C1',
    'trainee': 'E96E20',
}
ROLE_LABELS = {'chair': 'Chair', 'panellist': 'Panellist', 'trainee': 'Trainee'}
GRAY = 'B7B7B7'            # no feedback / not allocated that round
TITLE_FILL = '0B3C5D'      # banner behind the title
SCORE_FMT = '0.00'         # per-round scores: shown with 2 decimals (the cell holds 4)
TEST_FMT = 'General'       # judge test score = the exact Base Score from Tabbycat, no rounding
AVG_FMT = '0.000'          # overall averages (3 decimals, as in your sheets)
TOTAL_FMT = '0;-0;"-"'    # total feedback: shows "-" when it is 0

def solid(hex_color):
    """Solid fill with an explicit, fully-opaque colour (renders the same in Excel and Google Sheets)."""
    argb = 'FF' + hex_color.upper().lstrip('#')
    return PatternFill(fill_type='solid', start_color=argb, end_color=argb, fgColor=argb, bgColor=argb)


# Finished exports waiting to be downloaded (single gunicorn worker -> in-memory is fine)
EXPORT_CACHE = OrderedDict()
MAX_CACHED = 6


# =============================================================================
# SMALL HELPERS
# =============================================================================

def listify(value):
    if value is None:
        return []
    if isinstance(value, (list, tuple, set)):
        return list(value)
    return [value]


def url_id(value, kind):
    """
    Pull the numeric id out of an API hyperlink, e.g.
      url_id('https://x/api/v1/tournaments/t/adjudicators/42', 'adjudicators') -> 42
    Accepts plain ints / digit strings / {'url':..., 'id':...} dicts too.
    """
    if value is None:
        return None
    if isinstance(value, dict):
        if value.get('id') is not None:
            return url_id(value.get('id'), kind)
        return url_id(value.get('url'), kind)
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    text = str(value).strip().rstrip('/')
    if text.isdigit():
        return int(text)
    match = re.search(r'/%s/(\d+)' % re.escape(kind), text)
    return int(match.group(1)) if match else None


def to_float(value):
    if value is None or value == '':
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def parse_round_filter(text):
    """'1-4, 6' -> {1,2,3,4,6}. Blank -> None (meaning 'automatic')."""
    text = (text or '').strip()
    if not text:
        return None
    wanted = set()
    for part in re.split(r'[,\s;]+', text):
        if not part:
            continue
        if '-' in part:
            a, _, b = part.partition('-')
            if a.isdigit() and b.isdigit():
                lo, hi = sorted((int(a), int(b)))
                wanted.update(range(lo, hi + 1))
        elif part.isdigit():
            wanted.add(int(part))
    return wanted or None


# =============================================================================
# TABBYCAT API READER (GET only)
# =============================================================================

class ApiError(Exception):
    def __init__(self, status, url, body=''):
        self.status = status
        self.url = url
        self.body = body
        super().__init__(f"HTTP {status} on {url}")

    def friendly(self):
        short = self.url
        if '/api/' in short:
            short = '/api/' + short.split('/api/', 1)[1]
        if self.status == 401:
            return (f"401 Unauthorized on {short} - the API token is missing/invalid. "
                    f"Use a token from an ADMIN account (Change Password page).")
        if self.status == 403:
            return (f"403 Forbidden on {short} - this account has no access to that data. "
                    f"Feedback and pairings need tab/adjudication-core (admin) access.")
        if self.status == 404:
            return f"404 Not Found on {short} - check the URL and tournament slug."
        return f"HTTP {self.status} on {short}: {self.body[:150]}"


class TabbycatReader:
    def __init__(self, base_url, token, tournament_slug, username=None, password=None):
        self.base_url = base_url.strip().rstrip('/')
        self.token = (token or '').strip()
        self.slug = (tournament_slug or '').strip().strip('/')
        self.username = username
        self.password = password
        self.session = requests.Session()
        self.session.headers.update({
            'User-Agent': 'TabbycatExporter/1.0 (Render; Python requests)',
            'Accept': 'application/json',
        })
        self.auth_method = None
        self._authenticate()

    # ---- URLs --------------------------------------------------------------
    def global_url(self, path):
        return f"{self.base_url}/api/v1{path}"

    def tournament_url(self, path=''):
        return f"{self.base_url}/api/v1/tournaments/{self.slug}{path}"

    # ---- authentication (same idea as the importer) -------------------------
    def _authenticate(self):
        probe = self.tournament_url('/adjudicators')
        if self.token:
            self.session.headers['Authorization'] = f'Token {self.token}'
            if self._status(probe) == 200:
                self.auth_method = 'token'
                return

        if self.username and self.password:
            self.session.headers.pop('Authorization', None)
            if self._login_session() and self._status(probe) == 200:
                self.auth_method = 'session'
                return

        # Keep the token header so later errors show the real status code
        if self.token:
            self.session.headers['Authorization'] = f'Token {self.token}'

    def _status(self, url):
        try:
            return self.session.get(url, timeout=20).status_code
        except requests.exceptions.RequestException:
            return 0

    def _login_session(self):
        try:
            login_url = f"{self.base_url}/accounts/login/"
            resp = self.session.get(login_url, timeout=15)
            match = re.search(r'name="csrfmiddlewaretoken" value="([^"]+)"', resp.text)
            csrf = match.group(1) if match else ''
            self.session.post(login_url, data={
                'username': self.username,
                'password': self.password,
                'csrfmiddlewaretoken': csrf,
                'next': '/',
            }, headers={'Referer': login_url}, timeout=15)
            return self.session.get(f"{self.base_url}/database/", timeout=15).status_code == 200
        except Exception:
            return False

    # ---- reading ------------------------------------------------------------
    def get_json(self, url, retries=3):
        last_error = None
        for attempt in range(retries):
            try:
                resp = self.session.get(url, timeout=60)
            except requests.exceptions.RequestException as exc:
                last_error = exc
                continue
            if resp.status_code == 200:
                try:
                    return resp.json()
                except ValueError:
                    raise ApiError(resp.status_code, url, 'Response was not JSON')
            if resp.status_code == 429:
                import time
                time.sleep(2 ** attempt)
                continue
            raise ApiError(resp.status_code, url, resp.text)
        raise ApiError(0, url, str(last_error) if last_error else 'Request failed')

    def get_list(self, url):
        """GET a list endpoint; tolerates {'results': [...], 'next': url} pagination."""
        data = self.get_json(url)
        items = []
        guard = 0
        while True:
            if isinstance(data, list):
                items.extend(data)
                break
            if isinstance(data, dict) and isinstance(data.get('results'), list):
                items.extend(data['results'])
                nxt = data.get('next')
                guard += 1
                if nxt and guard < 200:
                    data = self.get_json(nxt)
                    continue
            break
        return items

    def test_connection(self):
        diag = {'ok': False, 'auth_method': self.auth_method, 'steps': [], 'suggestion': ''}

        def step(label, url, need_list=False):
            try:
                resp = self.session.get(url, timeout=30)
                info = {'step': label, 'status': resp.status_code, 'ok': resp.status_code == 200}
                if resp.status_code == 200 and need_list:
                    try:
                        data = resp.json()
                        if isinstance(data, dict) and 'results' in data:
                            data = data['results']
                        info['count'] = len(data) if isinstance(data, list) else None
                    except ValueError:
                        pass
                return info
            except requests.exceptions.RequestException as exc:
                return {'step': label, 'status': 0, 'ok': False, 'error': str(exc)}

        try:
            resp = self.session.get(self.base_url, timeout=10, allow_redirects=True)
            diag['steps'].append({'step': 'Base URL reachable', 'status': resp.status_code,
                                  'ok': resp.status_code < 500})
        except requests.exceptions.RequestException as exc:
            diag['steps'].append({'step': 'Base URL reachable', 'status': 0, 'ok': False, 'error': str(exc)})
            diag['suggestion'] = 'Cannot reach your Tabbycat URL. Check for typos.'
            return diag

        diag['steps'].append(step('Tournament (GET)', self.tournament_url()))
        diag['steps'].append(step('Rounds (GET)', self.tournament_url('/rounds'), True))
        adj = step('Adjudicators (GET, needs admin)', self.tournament_url('/adjudicators'), True)
        fb = step('Feedback (GET, needs admin)', self.tournament_url('/feedback'), True)
        diag['steps'].extend([adj, fb])

        if fb['ok'] and adj['ok']:
            diag['ok'] = True
            diag['suggestion'] = (f"Connection successful! Found {adj.get('count', '?')} adjudicators "
                                  f"and {fb.get('count', '?')} feedback records.")
        elif fb['status'] in (401, 403) or adj['status'] in (401, 403):
            diag['suggestion'] = ('Connected, but feedback/adjudicators are not accessible. Use the API token of '
                                  'an admin / adjudication-core account, or the session-login fallback.')
        elif fb['status'] == 404 or adj['status'] == 404:
            diag['suggestion'] = f'Tournament slug "{self.slug}" not found.'
        else:
            diag['suggestion'] = 'Unexpected response - see the steps below.'
        return diag


# =============================================================================
# FETCH RAW DATA
# =============================================================================

def fetch_raw(reader, opts):
    """Pull everything we need from the API (GET only)."""
    warnings = []

    tournament = {}
    try:
        tournament = reader.get_json(reader.tournament_url())
    except ApiError as exc:
        warnings.append(f"Could not read tournament name: {exc.friendly()}")
    tournament_name = (tournament.get('name') if isinstance(tournament, dict) else None) or reader.slug

    all_rounds = reader.get_list(reader.tournament_url('/rounds'))
    all_rounds = sorted([r for r in all_rounds if isinstance(r, dict) and r.get('seq') is not None],
                        key=lambda r: r['seq'])

    wanted = opts.get('round_filter')
    if wanted:
        rounds = [r for r in all_rounds if r['seq'] in wanted]
    elif opts.get('include_break_rounds'):
        rounds = list(all_rounds)
    else:
        rounds = [r for r in all_rounds if r.get('stage', 'P') == 'P']
    if not rounds:
        raise ValueError('No rounds matched your selection.')

    adjudicators = reader.get_list(reader.tournament_url('/adjudicators'))

    # Institutions: id -> {name, code}
    institutions = {}
    for url in (reader.global_url('/institutions'), reader.tournament_url('/institutions')):
        try:
            for inst in reader.get_list(url):
                inst_id = url_id(inst.get('id') if inst.get('id') is not None else inst.get('url'), 'institutions')
                if inst_id is not None:
                    institutions[inst_id] = {'name': inst.get('name') or '',
                                             'code': inst.get('code') or ''}
            if institutions:
                break
        except ApiError:
            continue
    if not institutions:
        warnings.append('Institutions could not be read; institution column may be blank.')

    # Pairings: who was chair / panellist / trainee in which round
    roles = {}      # (round_seq, adjudicator_id) -> role
    debates = {}    # debate_id -> {'round': seq, 'trainees': set(adjudicator_id)}
    for rnd in rounds:
        pairing_url = (rnd.get('_links') or {}).get('pairing') or \
            reader.tournament_url(f"/rounds/{rnd['seq']}/pairings")
        try:
            pairings = reader.get_list(pairing_url)
        except ApiError as exc:
            warnings.append(f"{rnd.get('name', 'Round ' + str(rnd['seq']))}: pairings unavailable "
                            f"({exc.friendly()}). Colours for this round may be missing.")
            continue
        for debate in pairings:
            if not isinstance(debate, dict):
                continue
            debate_id = debate.get('id')
            if debate_id is None:
                debate_id = url_id(debate.get('url'), 'pairings')
            panel = debate.get('adjudicators') or {}
            chair = url_id(panel.get('chair'), 'adjudicators') if isinstance(panel, dict) else None
            panellists = [url_id(x, 'adjudicators') for x in listify(panel.get('panellists'))] \
                if isinstance(panel, dict) else []
            trainees = [url_id(x, 'adjudicators') for x in listify(panel.get('trainees'))] \
                if isinstance(panel, dict) else []
            if chair is not None:
                roles[(rnd['seq'], chair)] = 'chair'
            for adj_id in panellists:
                if adj_id is not None:
                    roles[(rnd['seq'], adj_id)] = 'panellist'
            for adj_id in trainees:
                if adj_id is not None:
                    roles[(rnd['seq'], adj_id)] = 'trainee'
            if debate_id is not None:
                debates[debate_id] = {'round': rnd['seq'],
                                      'trainees': {t for t in trainees if t is not None}}

    feedback = reader.get_list(reader.tournament_url('/feedback'))

    return {
        'tournament_name': tournament_name,
        'rounds': rounds,
        'adjudicators': adjudicators,
        'institutions': institutions,
        'roles': roles,
        'debates': debates,
        'feedback': feedback,
        'warnings': warnings,
    }


# =============================================================================
# COMPUTE THE TAB
# =============================================================================

def compute_tab(raw, opts):
    rounds_all = raw['rounds']
    wanted_seqs = {r['seq'] for r in rounds_all}
    roles = raw['roles']
    debates = raw['debates']
    warnings = list(raw['warnings'])

    stats = {'feedback_total': len(raw['feedback']), 'feedback_used': 0, 'discarded': 0,
             'ignored': 0, 'trainee_source': 0, 'no_round': 0, 'no_score': 0,
             'outside_rounds': 0, 'unknown_role': 0}

    scores = defaultdict(list)       # (adj_id, round_seq) -> [scores]
    for fb in raw['feedback']:
        if not isinstance(fb, dict):
            continue
        if fb.get('confirmed') is False:
            stats['discarded'] += 1
            continue
        if fb.get('ignored'):
            stats['ignored'] += 1
            continue
        score = to_float(fb.get('score'))
        if score is None:
            stats['no_score'] += 1
            continue

        adj_id = url_id(fb.get('adjudicator'), 'adjudicators')
        if adj_id is None:
            continue

        debate_ref = fb.get('debate')
        debate_id = url_id(debate_ref, 'pairings')
        round_seq = None
        if debate_id in debates:
            round_seq = debates[debate_id]['round']
        if round_seq is None and debate_ref is not None:
            match = re.search(r'/rounds/(\d+)/', str(debate_ref))
            if match:
                round_seq = int(match.group(1))
        if round_seq is None:
            raw_round = fb.get('round_seq')
            round_seq = raw_round if isinstance(raw_round, int) else url_id(fb.get('round'), 'rounds')
        if round_seq is None:
            stats['no_round'] += 1
            continue
        if round_seq not in wanted_seqs:
            stats['outside_rounds'] += 1
            continue

        # Feedback written by a trainee does not count (Tabbycat rule)
        source = fb.get('source_adjudicator') or fb.get('source')
        src_text = str(source.get('url') if isinstance(source, dict) else source)
        if '/adjudicators/' in src_text:
            src_adj = url_id(source, 'adjudicators')
            if src_adj is not None and debate_id in debates and src_adj in debates[debate_id]['trainees']:
                stats['trainee_source'] += 1
                continue

        scores[(adj_id, round_seq)].append(score)
        stats['feedback_used'] += 1

    # Automatic round trimming: drop rounds that have no feedback at all
    if not opts.get('round_filter'):
        seqs_with_data = {seq for (_, seq) in scores}
        trimmed = [r for r in rounds_all if r['seq'] in seqs_with_data]
        rounds = trimmed or rounds_all
    else:
        rounds = rounds_all
    round_seqs = [r['seq'] for r in rounds]

    # Rows
    inst_mode = opts.get('institution_display', 'code')
    rows = []
    for adj in raw['adjudicators']:
        if not isinstance(adj, dict):
            continue
        adj_id = adj.get('id')
        if adj_id is None:
            adj_id = url_id(adj.get('url'), 'adjudicators')

        inst_id = url_id(adj.get('institution'), 'institutions')
        inst = raw['institutions'].get(inst_id) if inst_id is not None else None
        if inst:
            inst_text = (inst['code'] or inst['name']) if inst_mode == 'code' else (inst['name'] or inst['code'])
        else:
            inst_text = '—'

        cells = []
        total_count = 0
        total_sum = 0.0
        for seq in round_seqs:
            vals = scores.get((adj_id, seq), [])
            role = roles.get((seq, adj_id))
            if vals:
                avg = round(sum(vals) / len(vals), 4)      # the cell holds 4 decimals, the sheet shows 2
                total_count += len(vals)
                total_sum += avg * len(vals)               # same maths as the sheet's formula
                if role is None:
                    stats['unknown_role'] += 1
                cells.append({'avg': avg, 'count': len(vals), 'role': role})
            else:
                cells.append({'avg': None, 'count': None, 'role': role})

        base = to_float(adj.get('base_score'))
        overall = (total_sum / total_count) if total_count else base
        round_avgs = [c['avg'] for c in cells if c['avg'] is not None]
        # per-round average: the test score counts as one score when asked (as in the MPDC / UADC sheets)
        pool = ([base] if (opts.get('test_in_average') and opts.get('include_test_score') and base is not None) else []) + round_avgs
        round_average = (sum(pool) / len(pool)) if pool else base
        if len(pool) > 1:
            dropped = min(pool)
            drop_average = (sum(pool) - dropped) / (len(pool) - 1)
        else:
            dropped, drop_average = None, round_average

        rows.append({
            'id': adj_id,
            'name': adj.get('name') or '',
            'institution': inst_text,
            'breaking': bool(adj.get('breaking')),
            'core': bool(adj.get('adj_core')),
            'test_score': base,
            'cells': cells,
            'average': overall,
            'round_average': round_average,
            'drop_average': drop_average,
            'dropped': dropped,
            'rounds_judged': len(round_avgs),
            'total_feedback': total_count,
            'rank': None,
        })

    if opts.get('hide_empty'):
        rows = [r for r in rows if r['total_feedback'] > 0 or r['core']]

    # Ranking: core = CAP, everybody else with feedback ranked by average (ties share a rank)
    def metric(row):
        # the number the tab is ranked and sorted by
        if opts.get('export_mode') == 'rounds':
            return row['drop_average'] if opts.get('uadc') else row['round_average']
        return row['average']

    ranked = sorted([r for r in rows if not r['core'] and r['total_feedback'] > 0 and metric(r) is not None],
                    key=lambda r: (-round(metric(r), 3), r['name'].lower()))
    last_key, last_rank = None, 0
    for position, row in enumerate(ranked, start=1):
        key = round(metric(row), 3)
        if key != last_key:
            last_rank, last_key = position, key
        row['rank'] = last_rank
    for row in rows:
        if row['core']:
            row['rank'] = 'CAP'
        elif row['rank'] is None:
            row['rank'] = '-'

    def sort_key(row):
        # Default order: largest -> smallest "Average (number of Feedback)".
        # Adjudication core (CAP) are NOT pinned to the top; they sit wherever their average puts them.
        avg = metric(row) if metric(row) is not None else float('-inf')
        return (-avg, row['name'].lower())
    rows.sort(key=sort_key)
    for position, row in enumerate(rows, start=1):
        row['no'] = position

    if stats['unknown_role']:
        warnings.append(f"{stats['unknown_role']} round score(s) had no matching pairing, so they have no "
                        f"role colour. Check that the pairings of those rounds are accessible.")
    if stats['no_round']:
        warnings.append(f"{stats['no_round']} feedback record(s) could not be matched to a round and were skipped.")
    if not stats['feedback_total']:
        warnings.append('The feedback endpoint returned no records. Make sure the token belongs to an admin account.')

    return {
        'tournament_name': raw['tournament_name'],
        'mode': opts.get('export_mode', 'combined'),
        'uadc': bool(opts.get('uadc')),
        'rounds': rounds,
        'rows': rows,
        'stats': stats,
        'warnings': warnings,
    }


# =============================================================================
# EXCEL BUILDER
# =============================================================================

def build_workbook(tab, opts):
    include_test = opts.get('include_test_score', True)
    rounds = tab['rounds']
    rows = tab['rows']

    wb = Workbook()
    ws = wb.active
    ws.title = 'Final Tab'
    ws.sheet_view.showGridLines = False

    thin = Side(style='thin', color='000000')
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', indent=1)

    fixed = ['Final Rank', 'Name', 'Institution', 'Breaking'] + (['Judge Test Score'] if include_test else [])
    n_fixed = len(fixed)
    first_round_col = n_fixed + 1
    avg_rounds_col = first_round_col + 2 * len(rounds)   # average of the per-round averages
    avg_col = avg_rounds_col + 1                          # average weighted by number of feedback
    total_col = avg_col + 1
    last_col = total_col

    # Row 1: title banner
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    title = ws.cell(row=1, column=1, value=f"{tab['tournament_name']} Adjudicator Tab : Final Tab")
    title.font = Font(name='Arial', size=18, bold=True, color='FFFFFF')
    title.fill = solid(TITLE_FILL)
    title.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[1].height = 42

    # Row 2: colour legend
    legend = [('Chair', ROLE_COLORS['chair']), ('Panellist', ROLE_COLORS['panellist']),
              ('Trainee', ROLE_COLORS['trainee']), ('No feedback', GRAY)]
    ws.cell(row=2, column=1, value='Legend:').font = Font(name='Arial', size=9, bold=True)
    ws.cell(row=2, column=1).alignment = Alignment(horizontal='right', vertical='center')
    for i, (label, color) in enumerate(legend):
        cell = ws.cell(row=2, column=2 + i, value=label)
        cell.fill = solid(color)
        cell.font = Font(name='Arial', size=9, bold=True, color='FFFFFF')
        cell.alignment = center
        cell.border = box
    ws.row_dimensions[2].height = 18

    # Rows 4-5: headers
    head_font = Font(name='Arial', size=10, bold=True)
    for idx, label in enumerate(fixed, start=1):
        ws.merge_cells(start_row=4, start_column=idx, end_row=5, end_column=idx)
        cell = ws.cell(row=4, column=idx, value=label)
        cell.font, cell.alignment = head_font, center
        for r in (4, 5):
            ws.cell(row=r, column=idx).border = box
    for i, rnd in enumerate(rounds):
        c = first_round_col + 2 * i
        ws.merge_cells(start_row=4, start_column=c, end_row=4, end_column=c + 1)
        top = ws.cell(row=4, column=c, value=rnd.get('name') or f"Round {rnd['seq']}")
        top.font, top.alignment = head_font, center
        ws.cell(row=4, column=c).border = box
        ws.cell(row=4, column=c + 1).border = box
        for offset, label in enumerate(('Average Score', 'No. of Feedback')):
            sub = ws.cell(row=5, column=c + offset, value=label)
            sub.font, sub.alignment, sub.border = Font(name='Arial', size=9, bold=True), center, box
    for col, label in ((avg_rounds_col, 'Average (number of Rounds)'),
                       (avg_col, 'Average (number of Feedback)'),
                       (total_col, 'Total Number of Feedback')):
        ws.merge_cells(start_row=4, start_column=col, end_row=5, end_column=col)
        cell = ws.cell(row=4, column=col, value=label)
        cell.font, cell.alignment = head_font, center
        for r in (4, 5):
            ws.cell(row=r, column=col).border = box
    ws.row_dimensions[4].height = 24
    ws.row_dimensions[5].height = 30

    # Data rows
    body_font = Font(name='Arial', size=10)
    bold_font = Font(name='Arial', size=10, bold=True)
    white_bold = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    row_idx = 6
    for row in rows:
        col = 1
        rank_cell = ws.cell(row=row_idx, column=col, value=row['rank'])
        rank_cell.font, rank_cell.alignment, rank_cell.border = bold_font, center, box
        col += 1
        name_cell = ws.cell(row=row_idx, column=col, value=row['name'])
        name_cell.font, name_cell.alignment, name_cell.border = body_font, left, box
        col += 1
        inst_cell = ws.cell(row=row_idx, column=col, value=row['institution'])
        inst_cell.font, inst_cell.alignment, inst_cell.border = body_font, center, box
        col += 1
        brk = ws.cell(row=row_idx, column=col, value=bool(row['breaking']))   # TRUE/FALSE -> Insert > Checkbox
        brk.font, brk.alignment, brk.border = Font(name='Arial', size=10), center, box
        col += 1
        if include_test:
            test = ws.cell(row=row_idx, column=col, value=row['test_score'])
            test.number_format = TEST_FMT
            test.font, test.alignment, test.border = bold_font, center, box
            col += 1

        for cell_data in row['cells']:
            score_cell = ws.cell(row=row_idx, column=col)
            count_cell = ws.cell(row=row_idx, column=col + 1)
            if cell_data['avg'] is not None:
                score_cell.value = cell_data['avg']   # exact value; the 0.00 format shows 2 decimals
                score_cell.number_format = SCORE_FMT
                count_cell.value = cell_data['count']
                count_cell.number_format = '0'
                color = ROLE_COLORS.get(cell_data['role'])
                if color:
                    score_cell.fill = solid(color)
                    score_cell.font = white_bold
                else:
                    score_cell.font = bold_font
                count_cell.font = body_font
            else:
                for c in (score_cell, count_cell):
                    c.fill = solid(GRAY)
            for c in (score_cell, count_cell):
                c.alignment, c.border = center, box
            col += 2

        # ---- live formulas (so the numbers can be checked cell by cell) ----
        score_refs = [f"{get_column_letter(first_round_col + 2 * i)}{row_idx}" for i in range(len(rounds))]
        count_refs = [f"{get_column_letter(first_round_col + 2 * i + 1)}{row_idx}" for i in range(len(rounds))]
        total_ref = f"{get_column_letter(total_col)}{row_idx}"
        if include_test:
            test_ref = f"{get_column_letter(first_round_col - 1)}{row_idx}"
            fallback = f'IF({test_ref}="","",{test_ref})'      # no feedback yet -> show test score
        else:
            fallback = '""'

        if rounds:
            # Total Number of Feedback = sum of the per-round counts
            total_formula = "=SUM(" + ",".join(count_refs) + ")"
            # Average (number of Rounds) = plain average of the per-round averages
            avg_rounds_formula = f"=IFERROR(AVERAGE({','.join(score_refs)}),{fallback})"
            # Average (number of Feedback) = sum(score x count) / total feedback
            products = "+".join(f"N({s_})*N({c_})" for s_, c_ in zip(score_refs, count_refs))
            avg_formula = f"=IF({total_ref}=0,{fallback},({products})/{total_ref})"
        else:
            total_formula, avg_rounds_formula, avg_formula = 0, '', ''

        avg_rounds_cell = ws.cell(row=row_idx, column=avg_rounds_col, value=avg_rounds_formula)
        avg_rounds_cell.number_format = AVG_FMT
        avg_cell = ws.cell(row=row_idx, column=avg_col, value=avg_formula)
        avg_cell.number_format = AVG_FMT
        total = ws.cell(row=row_idx, column=total_col, value=total_formula)
        total.number_format = TOTAL_FMT
        for c in (avg_rounds_cell, avg_cell, total):
            c.font, c.alignment, c.border = bold_font, center, box
        ws.row_dimensions[row_idx].height = 18
        row_idx += 1

    # Filter / sort buttons on the header row (row 5). Use the dropdown on
    # "Average (number of Rounds)" or "Average (number of Feedback)" to sort
    # largest -> smallest or smallest -> largest. Rows move together, and the
    # formulas only reference their own row, so they stay correct after sorting.
    last_data_row = row_idx - 1
    if last_data_row >= 6:
        ws.auto_filter.ref = f"A5:{get_column_letter(last_col)}{last_data_row}"

    # Column widths
    widths = {1: 9, 2: 25.1, 3: 14.1, 4: 9.5}
    if include_test:
        widths[5] = 11
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    for i in range(len(rounds)):
        ws.column_dimensions[get_column_letter(first_round_col + 2 * i)].width = 10.5
        ws.column_dimensions[get_column_letter(first_round_col + 2 * i + 1)].width = 10.5
    ws.column_dimensions[get_column_letter(avg_rounds_col)].width = 13
    ws.column_dimensions[get_column_letter(avg_col)].width = 13
    ws.column_dimensions[get_column_letter(total_col)].width = 12

    ws.freeze_panes = 'C6'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def build_reference_workbook(tab, opts):
    """
    Sheets laid out like the reference files:
      mode 'feedback'        -> Davao sheet (score + number of feedback per round, weighted average, total, final rank)
      mode 'rounds'          -> MPDC sheet  (one average per round, final average, rank)
      mode 'rounds' + UADC   -> Bali sheet  (+ Average with Min. Score Dropped, No. of Rounds Judged, Dropped Score)
    Breaking is the 4th column; it holds TRUE/FALSE.
    """
    mode = tab['mode']
    uadc = tab['uadc'] and mode == 'rounds'
    include_test = opts.get('include_test_score', True)
    in_avg = include_test and opts.get('test_in_average', True)
    rounds, rows = tab['rounds'], tab['rows']
    n_rounds = len(rounds)
    per_round_cols = 2 if mode == 'feedback' else 1

    wb = Workbook()
    ws = wb.active
    ws.title = 'Final Tab'
    med = Side(style='medium', color='000000')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    arial = Font(name='Arial', size=10)
    bold = Font(name='Arial', size=10, bold=True)
    white_bold = Font(name='Arial', size=10, bold=True, color='FFFFFF')

    # ---- column plan (column A is a narrow spacer, the table starts in B) ----
    c = 2
    col_no = c; c += 1
    col_name = c; c += 1
    col_inst = c; c += 1
    col_brk = c; c += 1
    col_test = None
    if include_test:
        col_test = c; c += 1
    first_round = c
    c += per_round_cols * n_rounds
    if mode == 'feedback':
        col_avg, col_total, col_rank = c, c + 1, c + 2
        last_col = col_rank
    elif uadc:
        col_avg, col_drop_avg, col_judged, col_dropped = c, c + 1, c + 2, c + 3
        last_col = col_dropped
    else:
        col_avg, col_rank = c, c + 1
        last_col = col_rank
    if uadc:
        col_rank = col_no                      # Bali layout: "Final Rank" is the first column
    two_row_header = not uadc
    head_top = 3 if uadc else 2
    head_bottom = 3
    first_data = 4
    last_data = first_data + len(rows) - 1

    def put(r, col, value=None, font=arial, align=center, fmt=None, fill=None):
        cell = ws.cell(row=r, column=col, value=value)
        cell.font, cell.alignment = font, align
        if fmt:
            cell.number_format = fmt
        if fill:
            cell.fill = solid(fill)
        return cell

    def rcol(i, k=0):                           # column of round i (k = 1 -> the "No. of Feedback" column)
        return first_round + per_round_cols * i + k

    # ---- headers ----
    def fixed_header(col, label):
        if two_row_header:
            ws.merge_cells(start_row=head_top, start_column=col, end_row=head_bottom, end_column=col)
        put(head_top, col, label, bold)

    fixed_header(col_no, 'Final Rank' if uadc else 'No.')
    fixed_header(col_name, 'Name')
    fixed_header(col_inst, 'Institution')
    fixed_header(col_brk, 'Breaking')
    if include_test:
        fixed_header(col_test, 'Judge Test Score')
    if mode == 'feedback':
        for i, rnd in enumerate(rounds):
            ws.merge_cells(start_row=2, start_column=rcol(i), end_row=2, end_column=rcol(i, 1))
            put(2, rcol(i), rnd.get('name') or f"Round {rnd['seq']}", bold)
            put(3, rcol(i), 'Average Score', arial)
            put(3, rcol(i, 1), 'No. of Feedback', arial)
        fixed_header(col_avg, 'Average')
        fixed_header(col_total, 'Total Number of Feedback')
        fixed_header(col_rank, 'Final Rank (excluding CAP)')
    elif uadc:
        for i, rnd in enumerate(rounds):
            put(3, rcol(i), rnd.get('name') or f"Round {rnd['seq']}", bold)
        for col, label in ((col_avg, 'Average'), (col_drop_avg, 'Average with Min. Score Dropped'),
                           (col_judged, 'No. of Rounds Judged'), (col_dropped, 'Dropped Score')):
            put(3, col, label, bold)
    else:
        if n_rounds:
            ws.merge_cells(start_row=2, start_column=rcol(0), end_row=2, end_column=rcol(n_rounds - 1))
            put(2, rcol(0), 'Average Feedback Score', bold)
        for i, rnd in enumerate(rounds):
            put(3, rcol(i), rnd.get('name') or f"Round {rnd['seq']}", bold)
        fixed_header(col_avg, 'Final Average')
        fixed_header(col_rank, 'Rank')

    # ---- data ----
    for n, row in enumerate(rows):
        r = first_data + n
        rank_font = arial if (mode == 'rounds' and not uadc) else bold
        put(r, col_no, row['rank'] if uadc else row['no'], bold if uadc else arial)
        put(r, col_name, row['name'], arial, left)
        put(r, col_inst, row['institution'], arial)
        put(r, col_brk, bool(row['breaking']), arial)
        if include_test:
            put(r, col_test, row['test_score'], bold, fmt=TEST_FMT)
        for i, cd in enumerate(row['cells']):
            sc = ws.cell(row=r, column=rcol(i))
            sc.alignment = center
            ct = ws.cell(row=r, column=rcol(i, 1)) if mode == 'feedback' else None
            if ct is not None:
                ct.alignment, ct.font = center, arial
            if cd['avg'] is not None:
                sc.value = cd['avg']                       # 4 decimals in the cell, 2 shown
                sc.number_format = SCORE_FMT
                color = ROLE_COLORS.get(cd['role'])
                sc.font = white_bold if color else bold
                if color:
                    sc.fill = solid(color)
                if ct is not None:
                    ct.value, ct.number_format = cd['count'], '0'
            else:
                sc.font = white_bold
                sc.fill = solid(GRAY)
                if ct is not None:
                    ct.fill = solid(GRAY)

        L = get_column_letter
        score_refs = [f"{L(rcol(i))}{r}" for i in range(n_rounds)]
        if mode == 'feedback':
            count_refs = [f"{L(rcol(i, 1))}{r}" for i in range(n_rounds)]
            total_ref = f"{L(col_total)}{r}"
            fallback = f'IF({L(col_test)}{r}="","",{L(col_test)}{r})' if include_test else '""'
            if n_rounds:
                total_f = "=SUM(" + ",".join(count_refs) + ")"
                products = "+".join(f"N({a})*N({b})" for a, b in zip(score_refs, count_refs))
                avg_f = f"=IF({total_ref}=0,{fallback},({products})/{total_ref})"
            else:
                total_f, avg_f = 0, ''
            put(r, col_avg, avg_f, bold, fmt=AVG_FMT)
            put(r, col_total, total_f, arial, fmt=TOTAL_FMT)
            put(r, col_rank, row['rank'], bold)
        else:
            first_val = col_test if in_avg else first_round
            rng = f"{L(first_val)}{r}:{L(rcol(n_rounds - 1))}{r}" if n_rounds else None
            fb = f'IF({L(col_test)}{r}="","",{L(col_test)}{r})' if (include_test and not in_avg) else '""'
            avg_f = f"=IFERROR(AVERAGE({rng}),{fb})" if rng else ''
            put(r, col_avg, avg_f, bold if not uadc else arial, fmt=AVG_FMT)
            if uadc:
                rounds_rng = f"{L(rcol(0))}{r}:{L(rcol(n_rounds - 1))}{r}" if n_rounds else None
                drop_f = (f"=IF(COUNT({rng})>1,(SUM({rng})-MIN({rng}))/(COUNT({rng})-1),IFERROR(AVERAGE({rng}),\"\"))"
                          if rng else '')
                put(r, col_drop_avg, drop_f, bold, fmt=AVG_FMT)
                put(r, col_judged, f"=COUNT({rounds_rng})" if rounds_rng else 0, arial, fmt='0')
                put(r, col_dropped, f'=IF(COUNT({rng})>1,MIN({rng}),"")' if rng else '', arial, fmt=SCORE_FMT)
            else:
                put(r, col_rank, row['rank'], arial)
        ws.row_dimensions[r].height = 18

    # ---- medium outline around the header block and the data block (as in the reference sheets) ----
    def border_for(r, col, top_row, bottom_row):
        return Border(left=med if col == 2 else None, right=med if col == last_col else None,
                      top=med if r == top_row else None, bottom=med if r == bottom_row else None)
    for r in range(head_top, head_bottom + 1):
        for col in range(2, last_col + 1):
            ws.cell(row=r, column=col).border = border_for(r, col, head_top, head_bottom)
    for r in range(first_data, max(last_data, first_data) + 1):
        for col in range(2, last_col + 1):
            ws.cell(row=r, column=col).border = border_for(r, col, first_data, last_data)

    # ---- "Judge Allocations" colour key to the right of the table ----
    key_col = last_col + 2
    put(head_bottom, key_col, 'Judge Allocations:', bold, Alignment(horizontal='center', vertical='center'))
    for k, (label, role) in enumerate((('Chair', 'chair'), ('Panelist', 'panellist'), ('Trainee', 'trainee'))):
        put(first_data + k, key_col, label, white_bold, fill=ROLE_COLORS[role])

    # ---- widths / heights / filter / freeze ----
    ws.column_dimensions['A'].width = 2
    widths = {col_no: 10.1 if uadc else 6.4, col_name: 25.1, col_inst: 15.1, col_brk: 9.5}
    if include_test:
        widths[col_test] = 11
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    for i in range(n_rounds):
        for k in range(per_round_cols):
            ws.column_dimensions[get_column_letter(rcol(i, k))].width = 10 if mode == 'rounds' else 10.5
    for col in range(first_round + per_round_cols * n_rounds, last_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 13
    if mode == 'feedback':
        ws.column_dimensions[get_column_letter(col_total)].width = 12
        ws.column_dimensions[get_column_letter(col_rank)].width = 13
    if uadc:
        ws.column_dimensions[get_column_letter(col_drop_avg)].width = 14
    ws.column_dimensions[get_column_letter(last_col + 1)].width = 2.6
    ws.column_dimensions[get_column_letter(key_col)].width = 16.4
    ws.row_dimensions[1].height = 11.25
    ws.row_dimensions[head_top].height = 23.25 if two_row_header else 40
    if two_row_header:
        ws.row_dimensions[head_bottom].height = 27
    if last_data >= first_data:
        ws.auto_filter.ref = f"B{head_bottom}:{get_column_letter(last_col)}{last_data}"
    ws.freeze_panes = f"D{first_data}"
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# =============================================================================
# FLASK ROUTES
# =============================================================================

def reader_from_payload(data):
    return TabbycatReader(
        data.get('base_url', ''), data.get('token', ''), data.get('slug', ''),
        username=data.get('username') or None, password=data.get('password') or None)


@app.route('/')
def index():
    return render_template('index.html')


@app.route('/test-connection', methods=['POST'])
def test_connection():
    data = request.get_json() or {}
    if not data.get('base_url') or not data.get('slug'):
        return jsonify({'ok': False, 'steps': [], 'suggestion': 'Enter the Tabbycat URL and tournament slug first.'})
    try:
        return jsonify(reader_from_payload(data).test_connection())
    except Exception as exc:
        return jsonify({'ok': False, 'steps': [], 'suggestion': f'Error: {exc}'})


@app.route('/inspect', methods=['POST'])
def inspect():
    """Show one raw record from each endpoint so the field names can be checked."""
    data = request.get_json() or {}
    try:
        reader = reader_from_payload(data)
        out = {}
        rounds = reader.get_list(reader.tournament_url('/rounds'))
        out['rounds_count'] = len(rounds)
        first_pairing_url = None
        if rounds:
            first_pairing_url = (rounds[0].get('_links') or {}).get('pairing')

        def sample(label, url):
            try:
                items = reader.get_list(url)
                out[label] = {'count': len(items), 'first_record': items[0] if items else None}
            except ApiError as exc:
                out[label] = {'error': exc.friendly()}

        sample('adjudicators', reader.tournament_url('/adjudicators'))
        sample('feedback', reader.tournament_url('/feedback'))
        if first_pairing_url:
            sample('pairings_round_1', first_pairing_url)
        return jsonify(out)
    except Exception as exc:
        return jsonify({'error': str(exc)})


@app.route('/export', methods=['POST'])
def export():
    base_url = request.form.get('api_url', '').strip()
    token = request.form.get('api_token', '').strip()
    slug = request.form.get('tournament_slug', '').strip()
    username = request.form.get('api_username', '').strip() or None
    password = request.form.get('api_password', '').strip() or None

    if not all([base_url, token or (username and password), slug]):
        flash('Tabbycat URL, tournament slug and an API token (or admin login) are required.', 'error')
        return redirect(url_for('index'))

    opts = {
        'round_filter': parse_round_filter(request.form.get('rounds')),
        'include_break_rounds': request.form.get('include_break_rounds') == 'on',
        'institution_display': request.form.get('institution_display', 'code'),
        'include_test_score': request.form.get('include_test_score') == 'on',
        'hide_empty': request.form.get('hide_empty') == 'on',
        'export_mode': request.form.get('export_mode') if request.form.get('export_mode') in ('combined', 'feedback', 'rounds') else 'combined',
        'test_in_average': request.form.get('test_in_average') == 'on',
        'uadc_setting': request.form.get('uadc', 'auto'),
    }

    try:
        reader = TabbycatReader(base_url, token, slug, username=username, password=password)
        raw = fetch_raw(reader, opts)
        # UADC columns (Average with Min. Score Dropped, ...) only for the per-round tab of a UADC tournament
        name_hint = f"{raw.get('tournament_name', '')} {slug}".lower()
        opts['uadc'] = opts['export_mode'] == 'rounds' and (
            opts['uadc_setting'] == 'yes' or (opts['uadc_setting'] == 'auto' and 'uadc' in name_hint))
        tab = compute_tab(raw, opts)
        workbook = build_workbook(tab, opts) if opts['export_mode'] == 'combined' else build_reference_workbook(tab, opts)
    except ApiError as exc:
        flash(f'API error: {exc.friendly()}', 'error')
        return redirect(url_for('index'))
    except ValueError as exc:
        flash(str(exc), 'error')
        return redirect(url_for('index'))
    except Exception as exc:
        flash(f'Error: {exc}', 'error')
        return redirect(url_for('index'))

    token_id = uuid.uuid4().hex
    EXPORT_CACHE[token_id] = (f"{slug}_adjudicator_tab_{opts['export_mode']}.xlsx", workbook.getvalue())
    while len(EXPORT_CACHE) > MAX_CACHED:
        EXPORT_CACHE.popitem(last=False)

    return render_template('results.html', tab=tab, download_id=token_id,
                           role_colors=ROLE_COLORS, gray=GRAY, preview_rows=tab['rows'][:25],
                           include_test=opts['include_test_score'], mode=opts['export_mode'], uadc=opts['uadc'])


@app.route('/download/<download_id>')
def download(download_id):
    item = EXPORT_CACHE.get(download_id)
    if not item:
        flash('That export has expired. Please run the export again.', 'error')
        return redirect(url_for('index'))
    filename, content = item
    return send_file(io.BytesIO(content),
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=filename)


if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=False)
